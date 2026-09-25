"""导入导出：CSV / Excel / JSON / SQL INSERT 四种格式。"""

import csv
import datetime as dt
import json
from decimal import Decimal
from pathlib import Path

from db.errors import DbError

SUPPORTED_EXPORT = ("csv", "xlsx", "json", "sql")
SUPPORTED_IMPORT = ("csv", "xlsx", "json")


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", "ignore")
    return value


def export_rows(path, columns, rows, fmt="csv", table=None):
    path = Path(path)
    headers = [column[0] for column in columns]
    keys = [column[1] for column in columns]
    matrix = [[_plain(row.get(key)) for key in keys] for row in rows]
    fmt = fmt.lower()
    if fmt == "csv":
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(matrix)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = table or "导出结果"
        sheet.append(headers)
        header_fill = PatternFill("solid", fgColor="2563EB")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in matrix:
            sheet.append(row)
        for index, _ in enumerate(headers, start=1):
            width = max(len(str(headers[index - 1])), *(len(str(r[index - 1])) for r in matrix)) if matrix else len(str(headers[index - 1]))
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(40, max(10, width + 4))
        sheet.freeze_panes = "A2"
        workbook.save(path)
    elif fmt == "json":
        payload = [{header: value for header, value in zip(headers, row)} for row in matrix]
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    elif fmt == "sql":
        target = table or "export_table"
        with path.open("w", encoding="utf-8") as handle:
            handle.write(f"-- 由 edu_desktop 导出，目标表 {target}\n")
            for row in matrix:
                values = ", ".join("NULL" if value is None else _sql_literal(value) for value in row)
                handle.write(f"INSERT INTO {target} ({', '.join(headers)}) VALUES ({values});\n")
    else:
        raise DbError(-1, f"不支持的导出格式: {fmt}")
    return len(matrix)


def _sql_literal(value):
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def import_rows(path, limit=None):
    path = Path(path)
    suffix = path.suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_IMPORT:
        raise DbError(-1, f"不支持的导入格式: {suffix}")
    if suffix == "csv":
        rows = _read_csv(path, limit)
    elif suffix == "xlsx":
        rows = _read_xlsx(path, limit)
    else:
        rows = _read_json(path, limit)
    if not rows:
        raise DbError(-1, "文件没有可导入的数据行")
    headers = list(rows[0].keys())
    ordered = [{header: row.get(header) for header in headers} for row in rows]
    return headers, ordered


def _read_csv(path, limit):
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with path.open("r", newline="", encoding=encoding) as handle:
                reader = csv.DictReader(handle)
                rows = []
                for index, row in enumerate(reader):
                    if limit and index >= limit:
                        break
                    rows.append({key: (value if value != "" else None) for key, value in row.items()})
                return rows
        except UnicodeDecodeError:
            continue
    raise DbError(-1, "CSV 编码无法识别，请另存为 UTF-8")


def _read_xlsx(path, limit):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(name).strip() if name is not None else "" for name in next(iterator)]
    rows = []
    for index, values in enumerate(iterator):
        if limit and index >= limit:
            break
        rows.append({headers[i]: _plain(value) for i, value in enumerate(values) if i < len(headers)})
    workbook.close()
    return rows


def _read_json(path, limit):
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        for value in data.values():
            if isinstance(value, list):
                data = value
                break
    if not isinstance(data, list):
        raise DbError(-1, "JSON 结构必须是对象数组")
    return data[: limit or len(data)]
